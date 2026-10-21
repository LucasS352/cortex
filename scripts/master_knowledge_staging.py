"""Build a reviewable staging package for the ERP Master knowledge base.

This module is deliberately deterministic and does not write to PostgreSQL.
It inventories manuals, official MENU entries and extracted media, creates the
three initial pilot cards, and reports everything that still needs human
approval before ingestion as canonical knowledge.

Usage:
    python scripts/master_knowledge_staging.py
    python scripts/master_knowledge_staging.py --strict
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
CATALOG_PATH = ROOT / "output" / "master_screens" / "catalogo_oficial_telas_master.json"
MANUALS_DIR = ROOT / "Instrucoes_master"
ASSETS_DIR = ROOT / "assets" / "screens"
OUTPUT_DIR = ROOT / "output" / "master_screens"

STOPWORDS = {
    "a", "ao", "aos", "as", "com", "como", "da", "das", "de", "do", "dos",
    "e", "em", "na", "nas", "no", "nos", "o", "os", "para", "por", "um", "uma",
}


def _normalized(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.casefold())
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def _tokens(value: str) -> set[str]:
    return {part for part in _normalized(value).split() if len(part) > 1 and part not in STOPWORDS}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _relative(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def _image_dimensions(path: Path) -> tuple[int | None, int | None]:
    try:
        from PIL import Image

        with Image.open(path) as image:
            return image.size
    except Exception:
        return None, None


def _manual_text(path: Path) -> str:
    doc = Document(path)
    blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def _media(path: str, name: str, caption: str, after_step: int) -> dict[str, Any]:
    absolute = ROOT / path
    width, height = _image_dimensions(absolute) if absolute.exists() else (None, None)
    return {
        "id": name,
        "name": name,
        "caption": caption,
        "after_step": after_step,
        "path": path,
        "sha256": _sha256(absolute) if absolute.exists() else None,
        "width": width,
        "height": height,
    }


def load_catalog() -> list[dict[str, Any]]:
    return json.loads(CATALOG_PATH.read_text(encoding="utf-8"))


def build_pilot_cards() -> list[dict[str, Any]]:
    """Return explicit cards based only on evidence already reviewed in the project."""
    return [
        {
            "card_id": "master:produtos:importar-imagens",
            "status": "STAGING_REVIEW_REQUIRED",
            "title": "Importar imagens para o cadastro de produtos",
            "module": "Produtos e estoque",
            "source_manual": "Instrucoes_master/Fotos.docx",
            "menu": {
                "menu_id": None,
                "menu_group": "Utilitários",
                "screen_path": "Utilitários > Importação > Importar imagens",
                "form": None,
                "verification": "MANUAL_ONLY",
            },
            "intents": [
                "como importar imagens de produtos",
                "importar fotos no cadastro de produtos",
                "subir imagens em lote para produtos",
            ],
            "must_not_match": ["importar cadastro de produtos por planilha"],
            "shortcuts": [],
            "rules": [
                "O nome do arquivo deve conter o código do produto.",
                "Evitar caracteres especiais no nome do arquivo, como hífen e sublinhado, até homologação definitiva da regra.",
                "Repetir uma importação pode gerar duplicidade sem uma mensagem clara para o usuário.",
                "Confirmar se a empresa possui o serviço de armazenamento de imagens antes de orientar a importação.",
            ],
            "customer_answer": (
                "Para importar imagens dos produtos, salve cada arquivo com o código do produto no nome. "
                "Depois, acesse Utilitários > Importação > Importar imagens, selecione os arquivos e confira "
                "o critério de identificação antes de executar. Ao finalizar, valide o resultado na aba de "
                "fotos do cadastro do produto. Antes do procedimento, confirme se a empresa utiliza o serviço "
                "de armazenamento de imagens."
            ),
            "steps": [
                {"number": 1, "instruction": "Organize as imagens usando o código do produto no nome do arquivo."},
                {"number": 2, "instruction": "No Master, acesse Utilitários > Importação > Importar imagens."},
                {"number": 3, "instruction": "Selecione os arquivos e confira o critério usado para relacionar cada imagem ao produto."},
                {"number": 4, "instruction": "Execute a importação e valide as imagens no cadastro do produto."},
            ],
            "images": [
                _media(
                    "assets/screens/fotos/img_02.png",
                    "tela_importacao_imagens_produtos",
                    "Tela de importação de imagens e opções de identificação do produto.",
                    3,
                ),
                _media(
                    "assets/screens/fotos/img_01.png",
                    "aba_fotos_cadastro_produto",
                    "Aba de fotos do cadastro do produto usada para validar o resultado.",
                    4,
                ),
            ],
            "review_notes": [
                "O caminho não foi localizado no catálogo MENU; precisa de homologação visual.",
                "O print de e-mail img_00 foi excluído por conter dados identificáveis.",
            ],
        },
        {
            "card_id": "master:anymarket:transportadora-relacao-site",
            "status": "STAGING_REVIEW_REQUIRED",
            "title": "Relacionar transportadora da AnyMarket no Master",
            "module": "Integrações",
            "source_manual": None,
            "menu": {
                "menu_id": 93,
                "menu_group": "Cadastros",
                "screen_path": "Cadastros > Outros > Transportadoras > Relação - Site",
                "form": "FormCadastroTransportadora",
                "verification": "CATALOG_PLUS_CURATED_SCREENSHOT",
            },
            "intents": [
                "transportadora carrier não integrada na anymarket",
                "fazer de para da transportadora anymarket",
                "abrir relação site da transportadora com ctrl r",
            ],
            "must_not_match": [
                "produto sem ctrl r na anymarket",
                "de para de produtos anymarket",
            ],
            "shortcuts": [
                {"key": "Ctrl+R", "type": "CONTEXTUAL", "context": "Transportadora selecionada"}
            ],
            "rules": [
                "O relacionamento deve ser feito para a transportadora, não para o produto.",
                "O código externo deve corresponder ao retorno da plataforma, como Carrier quando exibido pela AnyMarket.",
            ],
            "customer_answer": (
                "A transportadora precisa ser relacionada com o código recebido da AnyMarket. No cadastro de "
                "Transportadoras, selecione a transportadora correspondente e pressione Ctrl+R. Na janela "
                "Relação - Site, selecione a AnyMarket, informe o código externo apresentado pela plataforma e "
                "confirme. Depois, tente novamente a integração do pedido."
            ),
            "steps": [
                {"number": 1, "instruction": "Abra o cadastro de Transportadoras e localize a transportadora correta."},
                {"number": 2, "instruction": "Com a transportadora selecionada, pressione Ctrl+R para abrir Relação - Site."},
                {"number": 3, "instruction": "Selecione a integração AnyMarket e informe o código externo retornado pela plataforma."},
                {"number": 4, "instruction": "Confirme o relacionamento e tente novamente a integração do pedido."},
            ],
            "images": [
                _media(
                    "assets/screens/anymarket_transportadora/ctrl_r_relacao_site.png",
                    "relacao_site_transportadora_anymarket",
                    "Janela Relação - Site usada no DE/PARA de transportadora da AnyMarket.",
                    3,
                )
            ],
            "review_notes": [
                "O atalho Ctrl+R é contextual e não aparece no campo MENU_ATALHO.",
                "A redação final deve ser homologada com um ticket real antes da publicação.",
            ],
        },
        {
            "card_id": "master:fiscal:nota-complementar-sem-difal",
            "status": "STAGING_REVIEW_REQUIRED",
            "title": "Emitir nota complementar sem cálculo de DIFAL",
            "module": "Fiscal e notas",
            "source_manual": "Instrucoes_master/Nota Complementar sem calculo de DIFAL.docx",
            "menu": {
                "menu_id": 1006,
                "menu_group": "Cadastros",
                "screen_path": "Cadastros > Fiscal > Tributação > Grupo de % de ICMS de produto por UF",
                "form": "FormCadastroGrupoICMSUF",
                "verification": "CATALOG_PLUS_MANUAL_PLUS_SCREENSHOT",
            },
            "intents": [
                "nota complementar sem cálculo de difal",
                "como deixar a alíquota de difal zerada",
                "não informar difal em nota complementar",
            ],
            "must_not_match": [
                "rejeição de difal na sefaz",
                "nota fiscal não calcula difal por erro",
            ],
            "shortcuts": [],
            "rules": [
                "Cadastrar uma regra específica com UF de origem, UF de destino, vigência e alíquota igual a zero.",
                "Vincular o grupo criado ao produto na aba Impostos.",
                "Recalcular a nota depois de vincular o grupo.",
                "Não aplicar esse procedimento a uma rejeição fiscal sem confirmar que a intenção é não calcular o DIFAL.",
            ],
            "customer_answer": (
                "Para emitir a nota complementar sem calcular DIFAL, cadastre uma regra em Cadastros > Fiscal > "
                "Tributação > Grupo de % de ICMS de produto por UF. Informe as UFs de origem e destino, a vigência "
                "e deixe a alíquota em 0%. Depois, vincule esse grupo ao produto na aba Impostos e recalcule a nota."
            ),
            "steps": [
                {"number": 1, "instruction": "Acesse Cadastros > Fiscal > Tributação > Grupo de % de ICMS de produto por UF."},
                {"number": 2, "instruction": "Informe UF de origem, UF de destino, data inicial e data final; deixe a alíquota em 0%."},
                {"number": 3, "instruction": "No cadastro do produto, abra a aba Impostos e vincule o grupo criado."},
                {"number": 4, "instruction": "Retorne à nota e execute novamente o cálculo."},
            ],
            "images": [
                _media(
                    "assets/screens/nota_complementar_sem_calculo_de_difal/img_02.png",
                    "menu_grupo_icms_produto_uf",
                    "Caminho no menu para o cadastro do grupo de ICMS por UF.",
                    1,
                ),
                _media(
                    "assets/screens/nota_complementar_sem_calculo_de_difal/img_01.png",
                    "grupo_icms_uf_aliquota_zero",
                    "Exemplo do grupo com origem, destino, vigência e alíquota zero.",
                    2,
                ),
                _media(
                    "assets/screens/nota_complementar_sem_calculo_de_difal/img_00.png",
                    "vinculo_grupo_icms_no_produto",
                    "Vínculo do grupo de ICMS por UF na aba Impostos do produto.",
                    3,
                ),
            ],
            "review_notes": ["Confirmar a abrangência por regime tributário e versão do Master."],
        },
    ]


@dataclass
class MatchCandidate:
    menu_id: int
    full_path: str
    form: str | None
    score: float


def _catalog_candidates(title: str, text: str, catalog: list[dict[str, Any]], limit: int = 5) -> list[dict[str, Any]]:
    title_tokens = _tokens(title)
    sample_tokens = _tokens(text[:5000])
    candidates: list[MatchCandidate] = []
    for entry in catalog:
        path_tokens = _tokens(entry.get("full_path") or "")
        if not path_tokens:
            continue
        title_overlap = len(title_tokens & path_tokens) / max(1, len(title_tokens | path_tokens))
        text_overlap = len(sample_tokens & path_tokens) / max(1, len(path_tokens))
        score = round((title_overlap * 0.75) + (text_overlap * 0.25), 4)
        if score > 0:
            candidates.append(MatchCandidate(entry["id"], entry["full_path"], entry.get("form"), score))
    candidates.sort(key=lambda item: (-item.score, item.full_path, item.menu_id))
    return [candidate.__dict__ for candidate in candidates[:limit]]


def build_match_candidates(catalog: list[dict[str, Any]]) -> list[dict[str, Any]]:
    asset_dirs = {_normalized(path.name): path for path in ASSETS_DIR.iterdir() if path.is_dir()}
    result = []
    for manual in sorted(MANUALS_DIR.rglob("*.docx")):
        if manual.name.startswith("~$"):
            continue
        try:
            text = _manual_text(manual)
        except Exception as exc:
            result.append({"manual": _relative(manual), "error": str(exc)})
            continue
        slug_key = _normalized(manual.stem)
        asset_dir = asset_dirs.get(slug_key)
        images = sorted(_relative(path) for path in asset_dir.iterdir() if path.is_file()) if asset_dir else []
        result.append(
            {
                "manual": _relative(manual),
                "title": manual.stem,
                "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                "characters": len(text),
                "asset_directory": _relative(asset_dir) if asset_dir else None,
                "images": images,
                "menu_candidates": _catalog_candidates(manual.stem, text, catalog),
                "auto_publish": False,
            }
        )
    return result


def validate(cards: list[dict[str, Any]], catalog: list[dict[str, Any]]) -> dict[str, Any]:
    catalog_by_id = {entry["id"]: entry for entry in catalog}
    full_paths: dict[str, list[int]] = defaultdict(list)
    for entry in catalog:
        full_paths[entry.get("full_path") or ""].append(entry["id"])
    duplicate_paths = {path: ids for path, ids in full_paths.items() if path and len(ids) > 1}

    card_ids = [card["card_id"] for card in cards]
    issues: list[dict[str, Any]] = []
    for duplicate, count in Counter(card_ids).items():
        if count > 1:
            issues.append({"severity": "ERROR", "card_id": duplicate, "code": "DUPLICATE_CARD_ID"})

    for card in cards:
        menu = card["menu"]
        menu_id = menu.get("menu_id")
        if card.get("status") != "APPROVED":
            issues.append({"severity": "REVIEW", "card_id": card["card_id"], "code": "PENDING_HUMAN_APPROVAL"})
        if not card.get("source_manual"):
            issues.append({"severity": "REVIEW", "card_id": card["card_id"], "code": "SOURCE_MANUAL_NOT_LINKED"})
        if menu_id is None:
            issues.append({"severity": "REVIEW", "card_id": card["card_id"], "code": "MENU_NOT_CATALOGED"})
        elif menu_id not in catalog_by_id:
            issues.append({"severity": "ERROR", "card_id": card["card_id"], "code": "MENU_ID_NOT_FOUND"})
        for image in card.get("images", []):
            if not (ROOT / image["path"]).is_file():
                issues.append({"severity": "ERROR", "card_id": card["card_id"], "code": "IMAGE_NOT_FOUND", "path": image["path"]})
            if not image.get("name") or not image.get("caption") or not image.get("after_step"):
                issues.append({"severity": "ERROR", "card_id": card["card_id"], "code": "IMAGE_MAPPING_INCOMPLETE", "path": image["path"]})
        step_numbers = {step["number"] for step in card.get("steps", [])}
        for image in card.get("images", []):
            if image.get("after_step") not in step_numbers:
                issues.append({"severity": "ERROR", "card_id": card["card_id"], "code": "IMAGE_STEP_NOT_FOUND", "path": image["path"]})

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "summary": {
            "catalog_entries": len(catalog),
            "duplicate_catalog_paths": len(duplicate_paths),
            "pilot_cards": len(cards),
            "errors": sum(issue["severity"] == "ERROR" for issue in issues),
            "review_items": sum(issue["severity"] == "REVIEW" for issue in issues),
            "ready_for_canonical_publish": sum(card["status"] == "APPROVED" for card in cards),
        },
        "duplicate_catalog_paths": duplicate_paths,
        "issues": issues,
    }


def write_outputs(strict: bool = False) -> dict[str, Any]:
    catalog = load_catalog()
    cards = build_pilot_cards()
    matches = build_match_candidates(catalog)
    report = validate(cards, catalog)
    report["summary"]["manuals_scanned"] = len(matches)
    report["summary"]["manuals_with_assets"] = sum(bool(item.get("images")) for item in matches)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "staging_cards.json").write_text(json.dumps(cards, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_DIR / "manual_menu_match_candidates.json").write_text(json.dumps(matches, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_DIR / "staging_validation_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    if strict and report["summary"]["errors"]:
        raise SystemExit(1)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera staging revisável da base oficial do ERP Master")
    parser.add_argument("--strict", action="store_true", help="Retorna erro se houver falha estrutural")
    args = parser.parse_args()
    report = write_outputs(strict=args.strict)
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
